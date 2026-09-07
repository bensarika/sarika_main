"""Render AST -> DOCX via python-docx.

Uses the built-in Word styles ("Title", "Heading 1..3", "List Bullet",
"Table Grid") so the file opens with a navigable heading tree in Word/Docs.
Inline **bold** / *italic* markers are converted to runs.
"""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Pt, RGBColor
from docx.text.paragraph import Paragraph as DocxParagraph

from ps_render.ast import BulletList, Document, Heading, Note, Paragraph, Table

_TOKEN = re.compile(r"(\*\*.+?\*\*|\*(?!\*).+?\*)")
_MUTED = RGBColor(0x6B, 0x72, 0x80)
_WARN = RGBColor(0xB4, 0x53, 0x09)
_ERR = RGBColor(0xB9, 0x1C, 0x1C)


def _runs(par: DocxParagraph, text: str, *, color: RGBColor | None = None, size: int | None = None) -> None:
    for tok in _TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            run = par.add_run(tok[2:-2])
            run.bold = True
        elif tok.startswith("*") and tok.endswith("*"):
            run = par.add_run(tok[1:-1])
            run.italic = True
        else:
            run = par.add_run(tok)
        if color is not None:
            run.font.color.rgb = color
        if size is not None:
            run.font.size = Pt(size)


def to_docx(doc: Document, path: Path) -> Path:
    d = DocxDocument()
    d.core_properties.title = doc.title

    sub = d.add_paragraph()
    _runs(sub, doc.subtitle, color=_MUTED)
    d.add_paragraph(doc.title, style="Title")
    tbl = d.add_table(rows=0, cols=2)
    tbl.style = "Table Grid"
    for label, value in doc.identifiers:
        cells = tbl.add_row().cells
        cells[0].text = label
        cells[1].text = value
    foot = d.add_paragraph()
    _runs(foot, doc.footer, color=_MUTED, size=9)
    d.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    for b in doc.blocks:
        if isinstance(b, Heading):
            d.add_heading(f"{b.number}  {b.text}", level=min(b.level, 3))
        elif isinstance(b, Paragraph):
            p = d.add_paragraph()
            _runs(p, b.text)
            if b.annotation in {"mismatch", "unbound"}:
                tag = "claim mismatch" if b.annotation == "mismatch" else "unverified claim"
                run = p.add_run(f"  [{tag}]")
                run.font.color.rgb = _ERR if b.annotation == "mismatch" else _WARN
                run.font.size = Pt(8)
        elif isinstance(b, BulletList):
            for item in b.items:
                p = d.add_paragraph(style="List Bullet")
                _runs(p, item)
        elif isinstance(b, Table):
            cap = d.add_paragraph()
            _runs(cap, b.caption, color=_MUTED, size=9)
            cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
            t = d.add_table(rows=1, cols=len(b.header))
            t.style = "Table Grid"
            for i, h in enumerate(b.header):
                cell = t.rows[0].cells[i]
                cell.text = ""
                _runs(cell.paragraphs[0], f"**{h}**", size=9)
            for row in b.rows:
                cells = t.add_row().cells
                for i in range(len(b.header)):
                    cells[i].text = ""
                    _runs(cells[i].paragraphs[0], row[i] if i < len(row) else "", size=9)
            for f in b.footnotes:
                fp = d.add_paragraph()
                _runs(fp, f, color=_MUTED, size=8)
        elif isinstance(b, Note):
            p = d.add_paragraph()
            _runs(p, b.text, color=_WARN if b.kind == "warning" else _MUTED, size=9)
    d.save(str(path))
    return path
