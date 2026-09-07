"""ps_render: AST build, deterministic LaTeX (golden), DOCX structure, optional Tectonic PDF."""

from __future__ import annotations

import json
from pathlib import Path
from typing import Any

import pytest
from docx import Document as DocxDocument

from ps_model.schema import empty_model, reference_dir
from ps_render import RenderMeta, build_ast, compile_pdf, tectonic_available, to_docx, to_latex
from ps_render.ast import Heading, Note, Paragraph, Table
from ps_render.latex import escape

GOLDEN = Path(__file__).parent / "golden" / "syn_protocol.tex"
META = RenderMeta(version_label="v-test", generated_at="2026-01-01")


@pytest.fixture
def syn() -> dict[str, Any]:
    with (reference_dir() / "protocol_examples.json").open(encoding="utf-8") as fh:
        raw = json.load(fh)
    return next(m for m in raw["models"] if m["protocol"]["id"] == "SYN.protocol.v1")


BLOCKS = [
    {
        "id": "b1",
        "section_id": "section.2",
        "subsection_id": "section.2.1",
        "order": 0,
        "kind": "paragraph",
        "text": "AD is **chronic** & relapsing.",
        "provenance": "author",
        "approval": "approved",
        "claims": [],
    },
    {
        "id": "b2",
        "section_id": "section.3",
        "subsection_id": "section.3.3",
        "order": 0,
        "kind": "paragraph",
        "text": "Assessed at Week 12.",
        "provenance": "author",
        "approval": "unreviewed",
        "claims": [
            {
                "id": "c1",
                "path": "endpoints[easi75].time.offset.value",
                "value": 12,
                "text": "Week 12",
                "state": "mismatch",
                "model_value": 16,
            }
        ],
    },
]


def test_ast_covers_all_15_sections_and_flags_missing(syn: dict[str, Any]) -> None:
    doc = build_ast(syn, BLOCKS, META)
    h1 = [b for b in doc.blocks if isinstance(b, Heading) and b.level == 1]
    assert len(h1) == 15
    assert any(isinstance(b, Table) for b in doc.blocks)  # generated views produce tables
    assert any(isinstance(b, Note) for b in doc.blocks)  # unwritten subsections are visible
    mismatch = [
        b for b in doc.blocks if isinstance(b, Paragraph) and b.annotation and "mismatch" in b.annotation.lower()
    ]
    assert mismatch, "mismatched claims must be annotated in output"


def test_empty_model_renders(syn: dict[str, Any]) -> None:
    m = empty_model(protocol_id="W", name="n", indication="AD")
    tex = to_latex(build_ast(m, [], META))
    assert "\\begin{document}" in tex and "Not yet written" in tex


def test_latex_escapes_specials() -> None:
    assert (
        escape("50% & $10 #1_x {y} ~ ^ \\")
        == r"50\% \& \$10 \#1\_x \{y\} \textasciitilde{} \textasciicircum{} \textbackslash{}"
    )


def test_latex_golden(syn: dict[str, Any]) -> None:
    tex = to_latex(build_ast(syn, BLOCKS, META))
    if not GOLDEN.exists():
        GOLDEN.parent.mkdir(exist_ok=True)
        GOLDEN.write_text(tex, encoding="utf-8")
    assert tex == GOLDEN.read_text(encoding="utf-8"), "LaTeX output drifted; review and refresh the golden file"


def test_docx_headings_and_tables(syn: dict[str, Any], tmp_path: Path) -> None:
    path = to_docx(build_ast(syn, BLOCKS, META), tmp_path / "p.docx")
    d = DocxDocument(str(path))
    h1 = [p for p in d.paragraphs if p.style is not None and p.style.name == "Heading 1"]
    assert len(h1) == 15
    assert d.tables
    assert any(r.bold for p in d.paragraphs for r in p.runs if "chronic" in r.text)


@pytest.mark.skipif(not tectonic_available(), reason="tectonic not installed")
def test_pdf_compiles(syn: dict[str, Any], tmp_path: Path) -> None:
    pdf = compile_pdf(to_latex(build_ast(syn, BLOCKS, META)), tmp_path)
    assert pdf.exists() and pdf.stat().st_size > 10_000
