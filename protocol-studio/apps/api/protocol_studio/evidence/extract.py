"""Deterministic first-pass extraction of an uploaded source.

Produces *candidates* with page numbers: identifiers (NCT, protocol id,
version/amendment labels), sentences that define the primary endpoint, and the
first eligibility criteria. It never decides what the study *is*; the reviewer
confirms in Upload Study → Review Extraction, which creates the StudyRecord.
Any model-assisted pass (later) runs on top of this and is labelled as such.

Supported: PDF (pypdf text layer), Markdown/plain text, Excel (openpyxl, first
sheet header + rows). A PDF with no text layer yields ``text_extractable: false``
rather than a guess.
"""

from __future__ import annotations

import hashlib
import io
import logging
import re
from typing import Any

from openpyxl import load_workbook
from pypdf import PdfReader

log = logging.getLogger(__name__)

MAX_PAGES = 250
MEDIA_TYPES = {
    "pdf": "application/pdf",
    "md": "text/markdown",
    "markdown": "text/markdown",
    "txt": "text/plain",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}

_WS = re.compile(r"\s+")
_NCT = re.compile(r"NCT\d{8}")
_EUDRA = re.compile(r"\b20\d{2}-\d{6}-\d{2}(?:-\d{2})?\b")
_VERSION = re.compile(r"\b(?:Protocol\s+)?(?:Amendment|Version)\s*[:#]?\s*\d+(?:\.\d+)?\b", re.IGNORECASE)
_PRIMARY = re.compile(
    r"[Pp]rimary\s+(?:efficacy\s+)?[Ee]ndpoint[s]?\s*[:•\-]?\s*(?:•\s*)?([A-Z(][^.]{12,240}\.)",
)
_PCT = re.compile(r"percent(?:age)? change|PCFB|% change", re.IGNORECASE)
_ABS = re.compile(r"\b(?:absolute )?change (?:in|from)\b", re.IGNORECASE)
_RESP = re.compile(r"\b(?:EASI|IGA|vIGA|NRS)[- ]?(?:\d{2}|0/1|≥\s?\d)", re.IGNORECASE)
_WEEK = re.compile(r"[Ww]eek\s*(\d{1,2})")
_INCLUSION = re.compile(r"Inclusion criteria(.{40,1400})", re.IGNORECASE | re.DOTALL)
_AGE = re.compile(
    r"(\d{2})\s*(?:–|-|to)\s*(\d{2})\s*years|(?:≥|>=|at least)\s*(\d{2})\s*years|(\d{2})\s*years (?:of age )?(?:or older|and older)"
)


def sha256_hex(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def media_type_for(filename: str) -> str:
    ext = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    return MEDIA_TYPES.get(ext, "application/octet-stream")


def _pages_pdf(data: bytes) -> tuple[list[str], int, bool]:
    reader = PdfReader(io.BytesIO(data))
    n = len(reader.pages)
    pages: list[str] = []
    for i in range(min(n, MAX_PAGES)):
        try:
            pages.append(_WS.sub(" ", reader.pages[i].extract_text() or ""))
        except Exception as e:
            log.info("pdf page %d extraction failed: %s", i + 1, e)
            pages.append("")
    extractable = sum(len(p) for p in pages) > 200
    return pages, n, extractable


def _pages_xlsx(data: bytes) -> tuple[list[str], dict[str, Any]]:
    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = [[("" if v is None else str(v)) for v in r] for r in ws.iter_rows(values_only=True)]
    header = rows[0] if rows else []
    body = rows[1:51]
    text = " ".join(" ".join(r) for r in rows)
    return [text], {"sheet": ws.title, "header": header, "rows": body, "row_count": max(len(rows) - 1, 0)}


def classify_endpoint(sentence: str) -> dict[str, Any]:
    tr = (
        "percent_change"
        if _PCT.search(sentence)
        else "responder"
        if _RESP.search(sentence)
        else "absolute_change"
        if _ABS.search(sentence)
        else "other"
    )
    wk = _WEEK.search(sentence)
    instrument = next((i for i in ("EASI", "IGA", "vIGA", "NRS", "SCORAD", "DLQI", "BSA") if i in sentence), "")
    return {"transformation": tr, "timepoint_weeks": int(wk.group(1)) if wk else None, "instrument": instrument}


def parse_age(text: str) -> dict[str, Any]:
    m = _AGE.search(text)
    if not m:
        return {}
    if m.group(1) and m.group(2):
        return {"min_age": int(m.group(1)), "max_age": int(m.group(2))}
    lo = m.group(3) or m.group(4)
    return {"min_age": int(lo)} if lo else {}


def _add_unique(into: list[str], found: list[str]) -> None:
    for x in found:
        if x not in into:
            into.append(x)


def extract(filename: str, data: bytes) -> dict[str, Any]:
    """Return the extraction dict stored on SourceRecord.extraction."""
    mt = media_type_for(filename)
    out: dict[str, Any] = {
        "media_type": mt,
        "sha256": sha256_hex(data),
        "bytes": len(data),
        "pages": None,
        "text_extractable": None,
        "identifiers": {"nct": [], "eudract": [], "versions": []},
        "primary_endpoint_candidates": [],
        "eligibility_candidates": [],
        "table": None,
        "warnings": [],
    }
    if mt == "application/pdf":
        pages, n, ok = _pages_pdf(data)
        out["pages"], out["text_extractable"] = n, ok
        if n > MAX_PAGES:
            out["warnings"].append(f"only the first {MAX_PAGES} of {n} pages were scanned")
        if not ok:
            out["warnings"].append("no machine-readable text layer; needs OCR or manual entry")
    elif mt == MEDIA_TYPES["xlsx"]:
        pages, table = _pages_xlsx(data)
        out["table"], out["text_extractable"] = table, True
    else:
        pages = [_WS.sub(" ", data.decode("utf-8", errors="replace"))]
        out["text_extractable"] = bool(pages[0].strip())

    ncts: list[str] = []
    eudra: list[str] = []
    versions: list[str] = []
    for pno, text in enumerate(pages, 1):
        _add_unique(ncts, _NCT.findall(text))
        _add_unique(eudra, _EUDRA.findall(text))
        _add_unique(versions, _VERSION.findall(text))
        if len(out["primary_endpoint_candidates"]) < 6:
            for m in _PRIMARY.finditer(text):
                s = m.group(1).strip()
                if len(s) < 20 or s.lower().startswith("analysis"):
                    continue
                out["primary_endpoint_candidates"].append({"page": pno, "text": s, **classify_endpoint(s)})
        if len(out["eligibility_candidates"]) < 3:
            m2 = _INCLUSION.search(text)
            if m2 and re.search(r"\byears?\b", m2.group(1)):
                body = m2.group(1).strip()
                out["eligibility_candidates"].append({"page": pno, "text": body[:1200], **parse_age(body)})
    out["identifiers"] = {"nct": ncts[:10], "eudract": eudra[:5], "versions": versions[:8]}
    return out


def draft_record(source_id: str, filename: str, extraction: dict[str, Any], indication: str) -> dict[str, Any]:
    """A best-effort StudyRecord *draft* from an extraction, for the reviewer to correct. Everything is 'unverified'."""
    ids = extraction.get("identifiers", {})
    nct = (ids.get("nct") or [""])[0]
    cands = extraction.get("primary_endpoint_candidates") or []
    elig = extraction.get("eligibility_candidates") or []
    doc = {
        "source_id": source_id,
        "kind": "protocol",
        "title": filename,
        "version_label": (ids.get("versions") or [""])[0],
        "sha256": extraction.get("sha256", ""),
        "bytes": extraction.get("bytes", 0),
        "pages": extraction.get("pages"),
        "text_extractable": extraction.get("text_extractable"),
    }
    endpoints = []
    if cands:
        c = cands[0]
        endpoints.append(
            {
                "id": "primary_candidate",
                "role": "primary",
                "label": c["text"].rstrip("."),
                "instrument": c.get("instrument", ""),
                "transformation": c.get("transformation", "other"),
                "timepoint_weeks": c.get("timepoint_weeks"),
                "provenance": {
                    "document": filename,
                    "source_id": source_id,
                    "page": c["page"],
                    "quote": c["text"],
                    "status": "unverified",
                },
            }
        )
    criteria = []
    if elig and (elig[0].get("min_age") is not None):
        e = elig[0]
        criteria.append(
            {
                "id": "age_candidate",
                "kind": "inclusion",
                "category": "age",
                "text": e["text"][:300],
                "min_age": e.get("min_age"),
                "max_age": e.get("max_age"),
                "provenance": {"document": filename, "source_id": source_id, "page": e["page"], "status": "unverified"},
            }
        )
    return {
        "record_version": "0.2.0",
        "id": "",
        "indication": indication,
        "title": filename.rsplit(".", 1)[0].replace("_", " "),
        "intervention": "",
        "registry_ids": [nct] if nct else [],
        "documents": [doc],
        "endpoints": endpoints,
        "criteria": criteria,
        "review_scope": "Draft from deterministic extraction; every value is unverified until a reviewer confirms it.",
    }
